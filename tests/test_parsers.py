"""合成データ（架空の医療機関）でパーサの挙動を確認するテスト。

実データの Excel/PDF のレイアウトは確定していないため、想定しうる複数の形
（2段組み・1段組み・都県別シート・罫線つき PDF・罫線なし PDF）で動くことを確かめる。
実行: python -m pytest -q tests
"""

import datetime as dt
import io
import zipfile

import openpyxl
import pytest

from scraper.parse_excel import parse_zip
from scraper.parse_grid import grid_to_records
from scraper.parse_pdf import parse_pdf
from scraper.regions import find_municipality, is_tokyo_address
from scraper.wareki import find_dates

ROWS = [
    # (項番, 番号, 名称, 〒住所, 電話, 開設者, 管理者, 点数表, 指定日, 期間終, 勤務, 診療科, 病床, 備考, 理由, 状態)
    ("1", "01,3834,7", "架空眼科クリニック", "〒101－0021\n千代田区外神田１－２－３　テストビル５Ｆ", "03-1234-5678",
     "山田　太郎", "山田　太郎", "医", "令和 8年10月 1日", "令和14年 9月30日", "常　勤:    2\n医 2( )", "眼", "", "", "新規", "現存"),
    ("2", "06,1234,5", "医療法人社団架空会\nサンプル内科", "〒182－0024\n調布市布田２－３－４", "042-123-4567",
     "医療法人社団架空会\n理事長　鈴木　一郎", "佐藤　花子", "医", "令和 8年10月 1日", "令和14年 9月30日",
     "常　勤:    1\n非常勤:    1\n医 1( 1)", "内　皮　アレ", "一般 19", "", "新規", "現存"),
]
RETRO = [
    ("1", "57,2231,0", "架空歯科医院", "〒190－0012\n立川市曙町１－１－１", "042-111-2222",
     "高橋　次郎", "高橋　次郎", "歯", "令和 8年 9月 1日", "令和14年 8月31日", "常　勤:    1\n歯 1( )", "歯　小歯",
     "", "旧機関コード：57,2230,1", "交代", "現存"),
]


def _two_row_sheet(ws):
    ws.append(["新規指定医療機関一覧表"])
    ws.append(["［令和 8年 9月 1日　から　令和 8年 9月30日　医科　指定分］"])
    ws.append(["項番", "医療機関番号", "医療機関名称", "医療機関所在地", "開設者氏名", "管理者氏名", "点数表",
               "指定年月日", "登録理由", "病床数", "備考"])
    ws.append(["", "", "", "電話番号", "", "", "", "指定期間終", "現存/廃止", "勤務医数／診療科名", ""])

    def put(rows):
        for r in rows:
            (no, code, name, addr, tel, f, m, t, d, end, staff, dept, beds, note, reason, st) = r
            ws.append([int(no), code, name, addr, f, m, t, d, reason, beds, note])
            ws.append(["", "", "", tel, "", "", "", end, st, f"{staff}\n{dept}", ""])

    put(ROWS)
    ws.append(["［令和 8年 9月 1日　から　令和 8年 9月30日　歯科　遡及指定分］"])
    ws.append(["項番", "医療機関番号", "医療機関名称", "医療機関所在地", "開設者氏名", "管理者氏名", "点数表",
               "指定年月日", "登録理由", "病床数", "備考"])
    ws.append(["", "", "", "電話番号", "", "", "", "指定期間終", "現存/廃止", "勤務医数／診療科名", ""])
    put(RETRO)


def _zip_of(files: dict[str, bytes]) -> bytes:
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w") as z:
        for n, b in files.items():
            z.writestr(n, b)
    return buf.getvalue()


def _wb_bytes(wb) -> bytes:
    b = io.BytesIO()
    wb.save(b)
    return b.getvalue()


def check(recs):
    assert len(recs) == 3, recs
    a, b, c = recs
    assert a["code"] == "01,3834,7"
    assert a["name"] == "架空眼科クリニック"
    assert "千代田区外神田1-2-3" in a["address_raw"]
    assert a["tel"] == "03-1234-5678"
    assert a["designated"] == "2026-10-01" and a["period_end"] == "2032-09-30"
    assert a["reason"] == "新規" and a["status"] == "現存"
    assert a["section"] == "指定分" and a["section_category"] == "医科"
    assert b["name"] == "医療法人社団架空会サンプル内科"
    assert "理事長" in b["founder"] and b["manager"] == "佐藤 花子"
    assert c["section"] == "遡及指定分" and c["old_code"] == "57,2230,1"
    assert c["reason"] == "交代"
    # raw に全列が残っている
    assert any("備考" in h for h, _ in c["raw"])


def test_excel_two_row_layout():
    wb = openpyxl.Workbook()
    _two_row_sheet(wb.active)
    z = _zip_of({"131shinki_tokyo_r0810.xlsx": _wb_bytes(wb), "141shinki_kanagawa_r0810.xlsx": b"dummy"})
    recs, diag = parse_zip(z)
    check(recs)
    assert diag["selected"][0]["file"].startswith("131")


def test_excel_one_row_layout_sheets_per_pref():
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "神奈川県"
    ws.append(["項番", "医療機関番号", "医療機関名称", "医療機関所在地", "電話番号", "開設者氏名", "管理者氏名",
               "点数表", "指定年月日", "指定期間始", "登録理由"])
    ws.append([1, "0112345", "横浜テスト", "〒222-0033 横浜市港北区新横浜1-1", "045-000-0000", "x", "y", "医",
               dt.datetime(2026, 10, 1), dt.datetime(2026, 10, 1), "新規"])
    ws2 = wb.create_sheet("東京都")
    ws2.append(["項番", "医療機関番号", "医療機関名称", "医療機関所在地", "電話番号", "開設者氏名", "管理者氏名",
                "点数表", "指定年月日", "指定期間始", "登録理由"])
    ws2.append([1, 1383470, "港テスト薬局", "〒105-0001 港区虎ノ門1-1-1", "03-0000-0000", "株式会社テスト",
                "田中 三郎", "調", dt.datetime(2026, 10, 1), dt.datetime(2026, 10, 1), "新規"])
    z = _zip_of({"shinki_r0810.xlsx": _wb_bytes(wb)})
    recs, diag = parse_zip(z)
    assert len(recs) == 1
    r = recs[0]
    assert r["code"] == "13,8347,0" and r["table_code"] == "調" and r["designated"] == "2026-10-01"


def _pdf(path, ruled: bool):
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.cidfonts import UnicodeCIDFont
    from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle
    from reportlab.lib.styles import ParagraphStyle

    pdfmetrics.registerFont(UnicodeCIDFont("HeiseiKakuGo-W5"))
    st = ParagraphStyle("j", fontName="HeiseiKakuGo-W5", fontSize=7, leading=9)
    head = ["項番", "医療機関番号", "医療機関名称", "医療機関所在地／電話番号", "開設者氏名", "管理者氏名", "点数表",
            "指定年月日／指定期間終", "登録理由／現存"]
    story = []

    def table(rows, label):
        story.append(Paragraph(f"新規指定医療機関一覧表　［令和 8年 9月 1日　から　令和 8年 9月30日　{label}］", st))
        story.append(Spacer(1, 6))
        data = [[Paragraph(h, st) for h in head]]
        for r in rows:
            (no, code, name, addr, tel, f, m, t, d, end, staff, dept, beds, note, reason, s) = r
            data.append([Paragraph(x.replace("\n", "<br/>"), st) for x in
                         [no, code, name, f"{addr}\n{tel}", f + (f"\n{note}" if note else ""), m, t, f"{d}\n{end}", f"{reason}\n{s}"]])
        tb = Table(data, colWidths=[25, 55, 110, 170, 110, 70, 30, 85, 50])
        if ruled:
            tb.setStyle(TableStyle([("GRID", (0, 0), (-1, -1), 0.5, "black")]))
        story.append(tb)
        story.append(Spacer(1, 20))

    table(ROWS, "医科　指定分")
    table(RETRO, "歯科　遡及指定分")
    SimpleDocTemplate(str(path), pagesize=landscape(A4)).build(story)


@pytest.mark.parametrize("ruled", [True, False])
def test_pdf(tmp_path, ruled):
    p = tmp_path / "t.pdf"
    _pdf(p, ruled)
    recs, diag = parse_pdf(str(p))
    assert len(recs) == 3, diag
    a = recs[0]
    assert a["code"] == "01,3834,7"
    assert a["designated"] == "2026-10-01"
    assert a["tel"] == "03-1234-5678"
    assert "外神田" in a["address_raw"]
    assert recs[2]["old_code"] == "57,2230,1"
    assert recs[2]["section"] == "遡及指定分"


def test_regions():
    assert find_municipality("〒105-0001 東京都港区虎ノ門1") == "港区"
    assert find_municipality("〒189-0001 東村山市秋津町1") == "東村山市"
    assert find_municipality("西多摩郡瑞穂町箱根ケ崎1") == "瑞穂町"
    assert find_municipality("横浜市港北区新横浜1") is None
    assert not is_tokyo_address("さいたま市北区宮原町1")


def test_dates():
    assert find_dates("令和 8年10月 1日\n令和14年 9月30日") == ["2026-10-01", "2032-09-30"]
    assert find_dates("R8.10.1") == ["2026-10-01"]
    assert find_dates("令和元年5月1日") == ["2019-05-01"]
