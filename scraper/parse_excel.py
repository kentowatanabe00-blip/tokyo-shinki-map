"""Excel データ（都県別 ZIP）から東京都分を取り出してレコード化する。"""

from __future__ import annotations

import io
import re
import zipfile

from .parse_grid import grid_to_records, is_header_row, norm
from .regions import is_tokyo_address

TOKYO_HINT = re.compile(r"東京|tokyo|^13\d|_13[_.]", re.I)


def _zip_name(info: zipfile.ZipInfo) -> str:
    name = info.filename
    if not info.flag_bits & 0x800:  # UTF-8 フラグなし → たいてい Shift_JIS
        try:
            name = name.encode("cp437").decode("cp932")
        except (UnicodeEncodeError, UnicodeDecodeError):
            pass
    return name


def _xlsx_sheets(data: bytes):
    import openpyxl

    wb = openpyxl.load_workbook(io.BytesIO(data), data_only=True, read_only=False)
    for ws in wb.worksheets:
        grid = [list(r) for r in ws.iter_rows(values_only=True)]
        # 見出し行（とその直下の段）だけ、横方向の結合セルを展開する
        header_rows = {i for i, r in enumerate(grid) if is_header_row(r)}
        header_rows |= {i + 1 for i in header_rows}
        for rng in ws.merged_cells.ranges:
            r0, r1, c0, c1 = rng.min_row - 1, rng.max_row - 1, rng.min_col - 1, rng.max_col - 1
            if c0 == c1:
                continue
            for r in range(r0, r1 + 1):
                if r in header_rows and r < len(grid):
                    v = grid[r0][c0] if c0 < len(grid[r0]) else None
                    for c in range(c0, c1 + 1):
                        if c < len(grid[r]) and grid[r][c] in (None, ""):
                            grid[r][c] = v
        yield ws.title, grid


def _xls_sheets(data: bytes):
    import datetime as dt

    import xlrd

    book = xlrd.open_workbook(file_contents=data)
    for sh in book.sheets():
        grid = []
        for r in range(sh.nrows):
            row = []
            for c in range(sh.ncols):
                cell = sh.cell(r, c)
                if cell.ctype == xlrd.XL_CELL_DATE:
                    row.append(dt.datetime(*xlrd.xldate_as_tuple(cell.value, book.datemode)))
                else:
                    row.append(cell.value)
            grid.append(row)
        yield sh.name, grid


def iter_workbooks(zip_bytes: bytes):
    with zipfile.ZipFile(io.BytesIO(zip_bytes)) as zf:
        for info in zf.infolist():
            if info.is_dir():
                continue
            name = _zip_name(info)
            lower = name.lower()
            if lower.endswith((".xlsx", ".xlsm")):
                yield name, "xlsx", zf.read(info)
            elif lower.endswith(".xls"):
                yield name, "xls", zf.read(info)
            elif lower.endswith(".zip"):
                yield from iter_workbooks(zf.read(info))


def parse_zip(zip_bytes: bytes, context: dict | None = None) -> tuple[list[dict], dict]:
    diag: dict = {"files": [], "selected": []}
    books = list(iter_workbooks(zip_bytes))
    diag["files"] = [n for n, _, _ in books]

    tokyo_books = [b for b in books if TOKYO_HINT.search(b[0].rsplit("/", 1)[-1])]
    candidates = tokyo_books or books
    records: list[dict] = []
    for name, kind, data in candidates:
        sheets = list(_xlsx_sheets(data) if kind == "xlsx" else _xls_sheets(data))
        tokyo_sheets = [s for s in sheets if re.search(r"東京", s[0])]
        use = tokyo_sheets if (not tokyo_books and tokyo_sheets) else sheets
        for title, grid in use:
            ctx = {**(context or {})}
            recs, d = grid_to_records(grid, ctx)
            diag["selected"].append({"file": name, "sheet": title, **d})
            records.extend(recs)

    # 東京都のファイル／シートを特定できなかった場合は住所で絞る
    if not tokyo_books:
        before = len(records)
        records = [r for r in records if is_tokyo_address(r.get("address_raw") or "")]
        diag["filtered_by_address"] = {"before": before, "after": len(records)}
    diag["records"] = len(records)
    return records, diag
